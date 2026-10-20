#!/usr/bin/env python3
"""Build Portofolio-Analisis-HR-Attrition.pptx (16:9, 10 slides, Bahasa Indonesia).
Redesign design-taste: tema T1 (marine blue, bright red, butter yellow),
varian Minimalist. Konten dan semua angka identik dengan build_ppt_legacy.py.
Bright red hanya dipakai untuk angka attrition kunci."""
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn

GROUND = RGBColor(0xF7, 0xF4, 0xEE)
INK = RGBColor(0x0B, 0x2A, 0x44)
MARINE = RGBColor(0x0E, 0x4D, 0x7A)
RED = RGBColor(0xE5, 0x32, 0x2D)
BUTTER = RGBColor(0xF5, 0xC5, 0x42)
MUTED = RGBColor(0x54, 0x62, 0x6F)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PALE = RGBColor(0xC9, 0xD6, 0xE2)
FONT = "Noto Sans"
CHARTS = "/home/hatch/workspace/hr-attrition/charts"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]


def solid_bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def rect(slide, left, top, width, height, color):
    sp = slide.shapes.add_shape(1, Inches(left), Inches(top), Inches(width), Inches(height))
    sp.fill.solid()
    sp.fill.fore_color.rgb = color
    sp.line.fill.background()
    sp.shadow.inherit = False
    return sp


def textbox(slide, left, top, width, height):
    tx = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tx.text_frame
    tf.word_wrap = True
    return tf


def add_par(tf, text, size, color, bold=False, italic=False, first=False, align=None,
            space_after=8, space_before=0):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.space_after = Pt(space_after)
    p.space_before = Pt(space_before)
    if align is not None:
        p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.name = FONT
    r.font.color.rgb = color
    return p


def add_bullet(tf, text, size=16, color=INK, first=False, space_after=12):
    p = add_par(tf, text, size, color, first=first, space_after=space_after)
    pPr = p._p.get_or_add_pPr()
    pPr.append(pPr.makeelement(qn("a:buChar"), {"char": "\u2022"}))
    pPr.set("marL", "171450")
    pPr.set("indent", "-171450")
    return p


def content_title(slide, text):
    tf = textbox(slide, 0.7, 0.5, 11.9, 0.85)
    add_par(tf, text, 32, INK, bold=True, first=True)


def add_chart(slide, name, left, top, max_w, max_h):
    path = f"{CHARTS}/{name}.png"
    iw, ih = Image.open(path).size
    scale = min(max_w / iw, max_h / ih)
    w, h = iw * scale, ih * scale
    slide.shapes.add_picture(path, Inches(left), Inches(top), Inches(w), Inches(h))
    return w, h


def add_caption(slide, text, left, top, width):
    tf = textbox(slide, left, top, width, 0.4)
    add_par(tf, text, 12, MUTED, italic=True, first=True, align=PP_ALIGN.CENTER)


def bullet_block(slide, bullets, left, top, width, size=16, space_after=12):
    tf = textbox(slide, left, top, width, 5.4)
    for i, b in enumerate(bullets):
        add_bullet(tf, b, size=size, first=(i == 0), space_after=space_after)
    return tf


# ---------- Slide 1: Cover (Asymmetric Split) ----------
s = prs.slides.add_slide(blank)
solid_bg(s, GROUND)
rect(s, 8.9, 0, 4.433, 7.5, MARINE)
rect(s, 8.9, 0, 4.433, 0.14, BUTTER)
tf = textbox(s, 9.35, 2.35, 3.65, 2.6)
add_par(tf, "16,1%", 66, BUTTER, bold=True, first=True, space_after=4)
add_par(tf, "attrition rate keseluruhan", 15, WHITE, space_after=2)
add_par(tf, "237 dari 1.470 karyawan resign", 13, PALE, space_after=0)
tf = textbox(s, 0.9, 1.55, 7.5, 2.3)
add_par(tf, "Analisis HR Attrition dengan SQL dan Python", 40, INK, bold=True, first=True)
tf = textbox(s, 0.9, 3.75, 7.3, 0.8)
add_par(tf, "Studi Kasus: IBM HR Analytics Employee Attrition Dataset", 18, MARINE, first=True)
tf = textbox(s, 0.9, 5.85, 7.3, 1.1)
add_par(tf, "BAGHIZ ASSHIDDIQIE ARASYA", 20, MARINE, bold=True, first=True, space_after=4)
add_par(tf, "Data Analyst Portfolio Project", 13, MUTED, space_after=0)

# ---------- Slide 2: Latar Belakang ----------
s = prs.slides.add_slide(blank)
solid_bg(s, GROUND)
content_title(s, "Latar Belakang")
bullet_block(s, [
    "Dataset IBM HR Analytics: 1.470 karyawan dengan 35 kolom data demografi, pekerjaan, kompensasi, dan kepuasan",
    "Pertanyaan utama: kenapa karyawan resign dan siapa yang paling berisiko pergi",
    "Attrition rate keseluruhan 16,1 persen atau 237 karyawan, angka yang cukup tinggi untuk diabaikan",
    "Tujuan analisis: menemukan faktor pendorong resign lalu menyusun rekomendasi retensi yang konkret",
], 0.7, 1.75, 7.3, size=16, space_after=16)
tf = textbox(s, 8.6, 1.8, 4.0, 1.4)
add_par(tf, "1.470", 38, MARINE, bold=True, first=True, space_after=0)
add_par(tf, "karyawan tercatat di dataset", 13, MUTED, space_after=0)
tf = textbox(s, 8.6, 3.35, 4.0, 1.4)
add_par(tf, "35", 38, MARINE, bold=True, first=True, space_after=0)
add_par(tf, "kolom data per karyawan", 13, MUTED, space_after=0)
tf = textbox(s, 8.6, 4.9, 4.0, 1.4)
add_par(tf, "237", 38, MARINE, bold=True, first=True, space_after=0)
add_par(tf, "karyawan resign dari keseluruhan", 13, MUTED, space_after=0)

# ---------- Slide 3: Data dan Metodologi ----------
s = prs.slides.add_slide(blank)
solid_bg(s, GROUND)
content_title(s, "Data dan Metodologi")
bullet_block(s, [
    "Data dimuat ke SQLite lalu dianalisis dengan query SQL asli: agregasi, grouping, dan perbandingan antar segmen",
    "Pandas dipakai untuk eksplorasi lanjutan dan pembuatan 12 chart visual",
    "Attrition didefinisikan sebagai kolom Attrition bernilai Yes; rate dihitung sebagai Yes dibagi total karyawan",
    "Segmen yang dianalisis: departemen, job role, lembur, band gaji, kelompok umur, jarak rumah, kepuasan, masa kerja, demografi, dan frekuensi dinas",
], 0.7, 1.75, 11.9, size=16, space_after=16)
rect(s, 0.7, 5.7, 11.93, 0.95, MARINE)
tf = textbox(s, 0.7, 5.95, 11.93, 0.6)
add_par(tf, "Alur kerja: file CSV dimuat ke SQLite, dianalisis dengan query SQL, dieksplorasi dengan pandas, lalu divisualisasi menjadi 12 chart",
        15, WHITE, first=True, align=PP_ALIGN.CENTER)

# ---------- Slide 4: Temuan 1, role berisiko ----------
s = prs.slides.add_slide(blank)
solid_bg(s, GROUND)
content_title(s, "Temuan 1: Role Paling Berisiko Resign")
tf = textbox(s, 0.7, 1.5, 5.5, 1.5)
add_par(tf, "39,8%", 62, RED, bold=True, first=True, space_after=2)
add_par(tf, "attrition Sales Representative", 14, INK, space_after=0)
bullet_block(s, [
    "Sales Representative mencatat attrition 39,8 persen, hampir 2,5 kali rata rata perusahaan",
    "Laboratory Technician 23,9 persen dan staf Human Resources 23,1 persen menyusul di posisi berikutnya",
    "Departemen Sales tertinggi 20,6 persen, lalu Human Resources 19,0 persen, Research and Development 13,8 persen",
    "Sebaliknya Research Director hanya 2,5 persen dan Manager 4,9 persen, sangat stabil",
], 0.7, 3.15, 5.5, size=15, space_after=10)
w, h = add_chart(s, "03_attrition_jobrole", 6.5, 1.7, 6.3, 3.9)
add_caption(s, "Attrition rate per job role", 6.5, 1.7 + h + 0.12, w)

# ---------- Slide 5: Temuan 2, lembur ----------
s = prs.slides.add_slide(blank)
solid_bg(s, GROUND)
content_title(s, "Temuan 2: Lembur Melipatgandakan Risiko")
w, h = add_chart(s, "04_attrition_overtime", 0.7, 1.85, 5.3, 3.8)
add_caption(s, "Attrition rate: lembur vs tidak lembur", 0.7, 1.85 + h + 0.12, w)
tf = textbox(s, 6.6, 1.6, 6.0, 1.35)
add_par(tf, "30,5%", 54, MARINE, bold=True, first=True, space_after=2)
add_par(tf, "attrition karyawan yang lembur", 14, INK, space_after=0)
bullet_block(s, [
    "Karyawan yang lembur resign 30,5 persen, dibanding yang tidak lembur hanya 10,4 persen",
    "Artinya karyawan lembur berisiko hampir 3 kali lipat untuk pergi",
    "Dari 416 karyawan lembur, 127 di antaranya resign, porsi yang sangat besar",
    "Ini sinyal kuat burnout, terutama di tim Sales dan laboratorium yang lemburnya tinggi",
], 6.6, 3.1, 6.0, size=16, space_after=12)

# ---------- Slide 6: Temuan 3, gaji umur masa kerja ----------
s = prs.slides.add_slide(blank)
solid_bg(s, GROUND)
content_title(s, "Temuan 3: Gaji Rendah, Muda, dan Baru")
bullet_block(s, [
    "Gaji di bawah 3.000 resign 28,6 persen, jauh di atas band gaji lain yang berkisar 9 sampai 13 persen",
    "Karyawan di bawah 30 tahun resign 27,9 persen, lalu turun bertahap sampai 9,7 persen di usia 40an",
    "Masa kerja 0 sampai 2 tahun resign 29,8 persen, lalu turun ke 8,1 persen setelah 11 tahun",
    "Pola jelas: karyawan muda bergaji rendah di dua tahun pertama adalah kelompok paling rentan",
], 0.7, 1.75, 5.8, size=16, space_after=14)
w, h = add_chart(s, "05_attrition_income_band", 6.75, 1.85, 6.0, 3.85)
add_caption(s, "Attrition rate per band gaji bulanan", 6.75, 1.85 + h + 0.12, w)

# ---------- Slide 7: Temuan 4, kepuasan ----------
s = prs.slides.add_slide(blank)
solid_bg(s, GROUND)
content_title(s, "Temuan 4: Kepuasan Rendah, Resign Tinggi")
tf = textbox(s, 0.7, 1.6, 5.9, 2.2)
add_bullet(tf, "Work life balance buruk mencatat attrition 31,3 persen, faktor tunggal tertinggi dalam analisis kepuasan", size=16, first=True, space_after=12)
add_bullet(tf, "Kepuasan lingkungan rendah 25,4 persen dibanding sangat tinggi 13,5 persen", size=16, space_after=0)
tf = textbox(s, 6.75, 1.6, 5.85, 2.2)
add_bullet(tf, "Kepuasan kerja rendah 22,8 persen dibanding sangat tinggi 11,3 persen", size=16, first=True, space_after=12)
add_bullet(tf, "Survei kepuasan rutin bisa menjadi sistem peringatan dini sebelum karyawan benar benar pergi", size=16, space_after=0)
w, h = add_chart(s, "10_attrition_worklife_balance", 0.7, 3.55, 5.9, 3.3)
add_caption(s, "Attrition rate per tingkat work life balance", 0.7, 3.55 + h + 0.1, w)
tf = textbox(s, 7.1, 4.0, 5.5, 1.9)
add_par(tf, "31,3%", 52, MARINE, bold=True, first=True, space_after=2)
add_par(tf, "attrition pada work life balance buruk", 14, INK, space_after=0)

# ---------- Slide 8: Temuan 5, jarak dinas status ----------
s = prs.slides.add_slide(blank)
solid_bg(s, GROUND)
content_title(s, "Temuan 5: Jarak, Dinas, dan Status")
bullet_block(s, [
    "Jarak rumah 15 km atau lebih resign 20,6 persen, naik bertahap dari 13,6 persen untuk jarak dekat",
    "Karyawan yang sering dinas resign 24,9 persen, dibanding tidak pernah dinas hanya 8,0 persen",
    "Status single resign 25,5 persen, dua kali lipat dibanding menikah 12,5 persen",
    "Stagnasi juga berbahaya: 6 tahun atau lebih tanpa promosi resign 16,3 persen, naik lagi setelah sempat turun",
], 0.7, 1.75, 5.6, size=16, space_after=14)
w, h = add_chart(s, "12_attrition_demografi_travel", 6.6, 1.65, 6.2, 4.3)
add_caption(s, "Attrition: demografi dan frekuensi dinas", 6.6, 1.65 + h + 0.12, w)

# ---------- Slide 9: Rekomendasi ----------
s = prs.slides.add_slide(blank)
solid_bg(s, GROUND)
content_title(s, "Rekomendasi Retensi")
rekomendasi = [
    "Bangun program onboarding 2 tahun pertama yang terstruktur: buddy, mentor, dan check in karier berkala",
    "Kelola lembur dengan ketat plus cuti pengganti, prioritaskan tim Sales dan laboratorium",
    "Tinjau ulang gaji, bonus, dan jalur promosi untuk Sales Representative dan Laboratory Technician",
    "Tawarkan fleksibilitas hybrid, tunjangan transportasi, dan rotasi jadwal dinas untuk yang jaraknya jauh",
    "Jadikan survei kepuasan sebagai pemicu intervensi dini per tim sebelum terlambat",
]
for i, teks in enumerate(rekomendasi):
    y = 1.75 + i * 1.02
    rect(s, 0.7, y + 0.09, 0.16, 0.16, MARINE)
    tf = textbox(s, 1.15, y, 11.4, 0.85)
    add_par(tf, teks, 17, INK, first=True, space_after=0)

# ---------- Slide 10: Penutup (Designed Close) ----------
s = prs.slides.add_slide(blank)
solid_bg(s, MARINE)
rect(s, 0.9, 0.85, 0.55, 0.13, BUTTER)
tf = textbox(s, 0.9, 1.15, 11.5, 0.9)
add_par(tf, "Kesimpulan", 34, WHITE, bold=True, first=True)
tf = textbox(s, 0.9, 2.25, 11.5, 2.2)
add_bullet(tf, "Lembur, gaji rendah, dan masa kerja pendek adalah tiga pendorong resign terbesar", size=17, color=WHITE, first=True, space_after=12)
add_bullet(tf, "Fokus retensi pada Sales Representative, karyawan muda, dan dua tahun pertama", size=17, color=WHITE, space_after=12)
add_bullet(tf, "Tech stack: SQL di SQLite, Python dengan pandas, matplotlib, dan seaborn", size=17, color=WHITE, space_after=0)
tf = textbox(s, 0.9, 5.15, 11.5, 0.85)
add_par(tf, "BAGHIZ ASSHIDDIQIE ARASYA", 30, BUTTER, bold=True, first=True)
tf = textbox(s, 0.9, 6.15, 11.5, 1.0)
add_par(tf, "Terbuka untuk peluang data analyst dan diskusi seputar people analytics.", 15, WHITE, first=True, space_after=4)
add_par(tf, "github.com/baghizash/ibm-hr-attrition-people-analytics", 15, PALE, space_after=0)

prs.save("/home/hatch/workspace/your_files/Portofolio-Analisis-HR-Attrition.pptx")
print("saved")
